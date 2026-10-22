from __future__ import annotations

import csv
import io
import json
import zipfile
from collections.abc import Callable
from typing import Any

import openpyxl
import pytest
from worker_fakes import MAX_PARTS, PART_ROWS, FakeRowSource, Harness, make_row

from exportkit.models import FailureCode, JobStatus

MakeHarness = Callable[..., Harness]


def csv_rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def test_single_file_when_rows_fit_in_one_part(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(PART_ROWS, area="billing"))
    export_id = h.submit()
    stages = h.run()
    assert [json.loads(b)["stage"] for b in stages] == ["export"]

    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED
    assert job.row_count == PART_ROWS
    assert job.part_count == 1
    assert job.file_name.endswith(f"_UTC_{export_id}.csv")
    assert h.keys() == [job.object_key]
    rows = csv_rows(h.body(job.object_key))
    assert rows[0] == [
        "event_id",
        "occurred_at",
        "actor_id",
        "area",
        "action",
        "entity_type",
        "entity_id",
    ]
    assert len(rows) - 1 == PART_ROWS
    assert [n.export_id for n in h.notifier.sent] == [export_id]
    assert job.notified_at is not None


def test_empty_export_still_produces_a_file(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource())
    export_id = h.submit()
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.row_count == 0
    assert len(csv_rows(h.body(job.object_key))) == 1  # header only


def test_multi_part_split_cursor_continuity_and_zip(make_harness: MakeHarness) -> None:
    total = 4 * PART_ROWS + 7  # 5 parts, the last one partial
    h = make_harness(FakeRowSource.with_rows(total))
    export_id = h.submit(area="all")
    stages = [json.loads(b) for b in h.run()]
    assert [s["stage"] for s in stages] == ["export", "part", "part", "part", "part", "part", "zip"]
    assert [s["part"] for s in stages if s["stage"] == "part"] == [1, 2, 3, 4, 5]

    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED
    assert (job.part_count, job.parts_done, job.row_count) == (5, 5, total)
    assert job.file_name.endswith(f"_UTC_{export_id}.zip")
    assert "_All_Areas_" in job.file_name
    assert h.keys() == [job.object_key]  # parts deleted after zipping

    # Each part continued strictly after the previous part's last row.
    limits = [limit for _, limit in h.rows.streams]
    assert limits == [PART_ROWS] * 5
    archive = zipfile.ZipFile(io.BytesIO(h.body(job.object_key)))
    names = archive.namelist()
    assert names == [job.file_name.replace(".zip", f"_part{n:02d}-of-05.csv") for n in range(1, 6)]
    assert all(info.compress_type == zipfile.ZIP_DEFLATED for info in archive.infolist())
    ids: list[int] = []
    for name in names:
        part = csv_rows(archive.read(name))
        assert part[0][0] == "event_id"
        ids.extend(int(row[0]) for row in part[1:])
    assert len(ids) == total
    assert len(set(ids)) == total  # no row twice
    expected_order = [
        r.event_id
        for r in sorted(h.rows.rows, key=lambda r: (r.occurred_at, r.event_id), reverse=True)
    ]
    assert ids == expected_order
    assert len(h.notifier.sent) == 1


def test_too_large_fails_without_writing(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(MAX_PARTS * PART_ROWS + 1))
    export_id = h.submit(area="all")
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.FAILED
    assert job.failure_code is FailureCode.TOO_LARGE
    assert h.keys() == []
    assert [n.status for n in h.notifier.sent] == [JobStatus.FAILED]


def test_exactly_max_parts_is_allowed(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(MAX_PARTS * PART_ROWS))
    export_id = h.submit(area="all")
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.part_count == MAX_PARTS


def test_every_message_delivered_twice_is_harmless(make_harness: MakeHarness) -> None:
    total = 2 * PART_ROWS + 3
    h = make_harness(FakeRowSource.with_rows(total))
    export_id = h.submit(area="all")
    h.run(duplicate=True)
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.row_count == total
    archive = zipfile.ZipFile(io.BytesIO(h.body(job.object_key)))
    assert sum(len(csv_rows(archive.read(n))) - 1 for n in archive.namelist()) == total
    assert len(h.notifier.sent) == 1  # one e-mail, even with duplicates


def test_redelivery_after_completion_is_a_noop(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(PART_ROWS + 1))
    export_id = h.submit(area="all")
    h.run()
    before = h.job(export_id)
    streams = len(h.rows.streams)
    for body in (
        {"export_id": export_id, "stage": "export"},
        {"export_id": export_id, "stage": "part", "part": 1, "cursor": None},
        {"export_id": export_id, "stage": "zip"},
    ):
        h.sqs.send_message(QueueUrl=h.queue_url, MessageBody=json.dumps(body))
    h.run()
    assert len(h.rows.streams) == streams  # nothing re-read
    after = h.job(export_id)
    assert after.completed_at == before.completed_at and after.status is JobStatus.COMPLETED
    assert len(h.notifier.sent) == 1


def test_lost_next_message_is_healed_by_redelivery(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(2 * PART_ROWS + 1))
    export_id = h.submit(area="all")
    h.run(steps=1)  # export stage -> part 1 queued
    h.run(steps=1)  # part 1 done -> part 2 queued
    h.drain()  # ... and that part 2 message is lost
    h.sqs.send_message(
        QueueUrl=h.queue_url,
        MessageBody=json.dumps({"export_id": export_id, "stage": "part", "part": 1}),
    )
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.row_count == 2 * PART_ROWS + 1


def test_high_water_mark_excludes_rows_added_during_export(make_harness: MakeHarness) -> None:
    rows = FakeRowSource.with_rows(2 * PART_ROWS + 5)
    h = make_harness(rows)
    export_id = h.submit(area="all")
    h.run(steps=1)  # plan: pins high-water and count
    pinned = h.job(export_id).high_water_event_id
    assert pinned == 2 * PART_ROWS + 5
    # New activity arrives while the parts are being written (newer and with higher ids).
    rows.rows.extend(make_row(i, minutes=10_000 + i) for i in range(1000, 1015))
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.row_count == 2 * PART_ROWS + 5
    archive = zipfile.ZipFile(io.BytesIO(h.body(job.object_key)))
    ids = [int(r[0]) for n in archive.namelist() for r in csv_rows(archive.read(n))[1:]]
    assert max(ids) <= pinned


def test_row_appearing_below_high_water_fails_rowset_changed(make_harness: MakeHarness) -> None:
    rows = FakeRowSource([make_row(i) for i in range(1, 2 * PART_ROWS + 5) if i != 7])
    h = make_harness(rows)
    export_id = h.submit(area="all")
    h.run(steps=1)
    # An insert that got id 7 commits late: it is below the pinned mark but was not counted.
    rows.rows.append(make_row(7))
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.FAILED
    assert job.failure_code is FailureCode.ROWSET_CHANGED


def test_filters_and_detailed_columns(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(60))
    export_id = h.submit(area="billing", mode="detailed", filters={"action": "update"})
    h.run()
    job = h.job(export_id)
    expected = [r for r in h.rows.rows if r.area == "billing" and r.action == "update"]
    rows = csv_rows(h.body(job.object_key))
    assert rows[0][-1] == "details"
    assert len(rows) - 1 == len(expected) == job.row_count
    assert all(r[3] == "billing" and r[4] == "update" for r in rows[1:])
    assert json.loads(rows[1][-1]) == [{"field": "status", "old": "draft", "new": "sent"}]


def test_xlsx_single_and_zip_members_stored(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(PART_ROWS - 1, area="orders"))
    single = h.submit(area="orders", format="xlsx", mode="detailed")
    h.run()
    job = h.job(single)
    sheet = openpyxl.load_workbook(io.BytesIO(h.body(job.object_key)), read_only=True).active
    values = list(sheet.iter_rows(values_only=True))
    assert values[0][0] == "event_id" and values[0][-1] == "details"
    assert len(values) - 1 == PART_ROWS - 1

    h.rows.rows = [make_row(i) for i in range(1, 2 * PART_ROWS + 2)]
    multi = h.submit(area="all", format="xlsx")
    h.run()
    job = h.job(multi)
    archive = zipfile.ZipFile(io.BytesIO(h.body(job.object_key)))
    assert all(info.compress_type == zipfile.ZIP_STORED for info in archive.infolist())
    total = 0
    for name in archive.namelist():
        book = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True)
        total += book.active.max_row - 1
    assert total == job.row_count == 2 * PART_ROWS + 1


def test_notification_failure_does_not_fail_the_job(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource.with_rows(3))
    h.notifier.fail = True
    export_id = h.submit()
    h.run()
    job = h.job(export_id)
    assert job.status is JobStatus.COMPLETED and job.notified_at is None


def test_part_before_plan_is_retried(make_harness: MakeHarness) -> None:
    from export_worker.handler import process_batch

    h = make_harness(FakeRowSource.with_rows(3))
    export_id = h.submit()
    h.drain()
    record: dict[str, Any] = {
        "messageId": "m1",
        "body": json.dumps({"export_id": export_id, "stage": "part", "part": 1}),
        "attributes": {"ApproximateReceiveCount": "1"},
    }
    result = process_batch(h.worker, {"Records": [record]}, max_receive_count=3)
    assert result == {"batchItemFailures": [{"itemIdentifier": "m1"}]}
    assert h.job(export_id).status is JobStatus.QUEUED  # not the last attempt: left alone

    record["attributes"]["ApproximateReceiveCount"] = "3"
    process_batch(h.worker, {"Records": [record]}, max_receive_count=3)
    job = h.job(export_id)
    assert job.status is JobStatus.FAILED and job.failure_code is FailureCode.WORKER_ERROR
    assert [n.status for n in h.notifier.sent] == [JobStatus.FAILED]


@pytest.mark.parametrize("body", ["not json", '{"export_id":"x","stage":"nope"}'])
def test_malformed_messages_are_reported(make_harness: MakeHarness, body: str) -> None:
    from export_worker.handler import process_batch

    h = make_harness(FakeRowSource())
    record = {"messageId": "bad", "body": body, "attributes": {"ApproximateReceiveCount": "3"}}
    result = process_batch(h.worker, {"Records": [record]}, max_receive_count=3)
    assert result == {"batchItemFailures": [{"itemIdentifier": "bad"}]}


def test_unknown_job_is_dropped(make_harness: MakeHarness) -> None:
    h = make_harness(FakeRowSource())
    h.sqs.send_message(
        QueueUrl=h.queue_url,
        MessageBody=json.dumps({"export_id": "01JA2B3C4D5E6F7G8H9JKMNPQR", "stage": "export"}),
    )
    assert len(h.run()) == 1
